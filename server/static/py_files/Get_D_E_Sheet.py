import os
import Drawing
import openpyxl
from dialog import Dialog
from dialog.checkbox import CheckButtons
from dialog.choose_file import ChooseFile

Dlg0 = Dialog("Export Sheets to Excel")
CheckButtons(
    Dlg0, "sheet_types",
    ("Detail Sheets (D)", "Erection Sheets (E)"),
    label="Select sheet types to export:",
    default=["Detail Sheets (D)", "Erection Sheets (E)"]
)
ChooseFile(
    Dlg0, "output_path",
    label="Save Excel file to:",
    save=True,
    default="sheets_output.xlsx"
)

if not Dlg0.Run():
    print("Cancelled.")
else:
    selected = Dlg0.sheet_types
    include_d = "Detail Sheets (D)" in selected
    include_e = "Erection Sheets (E)" in selected
    output_path = Dlg0.output_path

    if not include_d and not include_e:
        print("No sheet type selected.")
    elif not output_path:
        print("No save path specified.")
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheets"

        col = 1
        d_col = e_col = None

        if include_d:
            ws.cell(row=1, column=col, value="Detail Sheets")
            d_col = col
            col += 1

        if include_e:
            ws.cell(row=1, column=col, value="Erection Sheets")
            e_col = col

        if include_d:
            detail_sheets = Drawing.GetAllDetailSheetNames()
            for i, s in enumerate(detail_sheets, start=2):
                ws.cell(row=i, column=d_col, value=Drawing.Drawing("Detail Sheet", s).drawing_name)

        if include_e:
            erection_sheets = Drawing.GetAllErectionSheetNames()
            for i, s in enumerate(erection_sheets, start=2):
                ws.cell(row=i, column=e_col, value=Drawing.Drawing("Erection Sheet", s).drawing_name)

        wb.save(output_path)
        os.startfile(output_path)
